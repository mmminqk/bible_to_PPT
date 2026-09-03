"""
pptx_generator.generator

PowerPoint 슬라이드 서식 적용, 도형 자동 매핑, 텍스트 렌더링,
블랙 슬라이드 삽입 및 최종 PPTX 생성을 담당하는 모듈.
"""

import os
import sys
import copy
import re
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

from constants import (
    SUPERSCRIPT_MAP,
    DEFAULT_BOLD_FONT,
    EMPHASIS_BOLD,
    EMPHASIS_UNDERLINE,
    get_full_style,
)
from loader import resource_path


# ─── 유틸리티 ────────────────────────────────────────────────────────────────
def hex_to_rgb(hex_color):
    """'1F3337' 또는 '#1F3337' 형식의 16진수 색상 코드를 RGBColor 객체로 변환."""
    h = hex_color.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _apply_run_style(run, font_name, font_size, color_hex, underline=False):
    """pptx run 요소에 폰트 이름, 크기, 색상, 밑줄 서식을 적용."""
    run.font.name      = font_name
    run.font.size      = Pt(font_size)
    run.font.color.rgb = hex_to_rgb(color_hex)
    run.font.underline = underline


def _superscript(token):
    """숫자 및 콜론 문자열을 유니코드 상첨자(⁰¹²³⁴⁵⁶⁷⁸⁹˸)로 변환."""
    return token.translate(SUPERSCRIPT_MAP)


def _verse_body_text(line):
    """앞선 절 번호 토큰을 제거하고 본문만 반환."""
    return ' '.join(line.split(' ')[1:])


def duplicate_slide_with_blank_layout(prs, slide):
    """기존 슬라이드의 shape들을 빈 레이아웃(slide_layouts[6]) 슬라이드에 복제하여 추가."""
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)
    for shape in slide.shapes:
        new_slide.shapes._spTree.insert_element_before(
            copy.deepcopy(shape.element), 'p:extLst'
        )
    return new_slide


def move_slide(prs, old_index, new_index):
    """Presentation 객체 내부에서 슬라이드를 old_index에서 new_index로 이동."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    elem = slides[old_index]
    xml_slides.remove(elem)
    if new_index >= len(xml_slides):
        xml_slides.append(elem)
    else:
        xml_slides.insert(new_index, elem)


def insert_black_slides(prs, group_sizes):
    """
    각 번호 항목의 마지막 슬라이드 뒤에 검은 슬라이드를 삽입한다.
    group_sizes: 각 번호 항목이 차지하는 슬라이드 수의 리스트 [n1, n2, ...]
    """
    insert_positions = []
    cumulative = 0
    for size in group_sizes:
        cumulative += size
        insert_positions.append(cumulative)

    for pos in reversed(insert_positions):
        blank_layout = prs.slide_layouts[6]
        new_slide = prs.slides.add_slide(blank_layout)
        bg = new_slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0, 0, 0)
        last_idx = len(prs.slides) - 1
        move_slide(prs, last_idx, pos)


# ─── 강조 구간 분할 ──────────────────────────────────────────────────────────
def _split_by_emphases(text, emphases):
    """
    text를 emphases 목록 기준으로 분할해
    [(segment_text, kind_or_None), ...] 리스트로 반환.
    kind_or_None이 None이면 일반 텍스트, 'bold'/'underline'이면 강조 구간.
    """
    if not emphases:
        return [(text, None)]

    hits = sorted(
        ((text.find(emp['text']), emp['text'], emp['kind'])
         for emp in emphases if text.find(emp['text']) != -1),
        key=lambda x: x[0],
    )

    segments = []
    cursor = 0
    for start, emph_text, kind in hits:
        if start < cursor:
            continue
        if start > cursor:
            segments.append((text[cursor:start], None))
        segments.append((emph_text, kind))
        cursor = start + len(emph_text)

    if cursor < len(text):
        segments.append((text[cursor:], None))

    return segments if segments else [(text, None)]


# ─── run 단위 텍스트 쓰기 ───────────────────────────────────────────────────
def _write_paragraph_with_emphasis(p, text, base_font, base_size, base_color, emphases, bold_font):
    """
    단락 p에 text를 강조 구간에 따라 여러 run으로 분할해 기록.
    - 일반 구간 : base_font / base_size / base_color
    - 굵게 구간 : bold_font / base_size / base_color
    - 밑줄 구간 : base_font / base_size / base_color + underline=True
    """
    segments = _split_by_emphases(text, emphases)
    for seg_text, kind in segments:
        run = p.add_run()
        run.text = seg_text
        if kind == EMPHASIS_BOLD:
            _apply_run_style(run, bold_font, base_size, base_color)
        elif kind == EMPHASIS_UNDERLINE:
            _apply_run_style(run, base_font, base_size, base_color, underline=True)
        else:
            _apply_run_style(run, base_font, base_size, base_color)


def _write_paragraph_plain(p, text, font, size, color, use_run=True):
    """강조 없이 단락 하나를 기록 (제목 등 단순 텍스트용)."""
    if use_run:
        run = p.add_run()
        run.text = text
        _apply_run_style(run, font, size, color)
    else:
        p.text = text
        p.font.name = font
        p.font.size = Pt(size)
        p.font.color.rgb = hex_to_rgb(color)


def _set_text_lines(tf, lines, font, size, color, use_run=True):
    """tf를 초기화한 뒤 lines를 강조 없이 한 줄씩 추가."""
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _write_paragraph_plain(p, line, font, size, color, use_run=use_run)


# ─── 템플릿 도형 자동 매핑 ──────────────────────────────────────────────────
_TAG_PATTERNS = {
    'kor_title': [
        r'\{\{\s*(?:한글_?제목|한글_?주소|kor_?title|kor_?addr)\s*\}\}',
        r'\{\s*(?:한글_?제목|한글_?주소|kor_?title|kor_?addr)\s*\}',
        r'구절\s*주소\s*\(\s*한\s*\)',
    ],
    'kor_body': [
        r'\{\{\s*(?:한글_?본문|한글_?말씀|개역개정|kor_?body|kor_?text)\s*\}\}',
        r'\{\s*(?:한글_?본문|한글_?말씀|개역개정|kor_?body|kor_?text)\s*\}',
        r'개역개정\s*본문',
    ],
    'eng_title': [
        r'\{\{\s*(?:영어_?제목|영어_?주소|eng_?title|eng_?addr)\s*\}\}',
        r'\{\s*(?:영어_?제목|영어_?주소|eng_?title|eng_?addr)\s*\}',
        r'구절\s*주소\s*\(\s*영\s*\)',
    ],
    'eng_body': [
        r'\{\{\s*(?:영어_?본문|영어_?말씀|esv|eng_?body|eng_?text)\s*\}\}',
        r'\{\s*(?:영어_?본문|영어_?말씀|esv|eng_?body|eng_?text)\s*\}',
        r'영어\s*본문',
        r'^abc$',
    ],
}

_NAME_KEYWORDS = {
    'kor_title': ['kor_title', 'kor_addr', '한글제목', '한글주소', '한글_제목'],
    'kor_body':  ['kor_body', 'kor_text', '한글본문', '한글_본문', '개역개정'],
    'eng_title': ['eng_title', 'eng_addr', '영어제목', '영어주소', '영어_제목'],
    'eng_body':  ['eng_body', 'eng_text', '영어본문', '영어_본문', 'esv'],
}


def detect_shape_mapping(slide):
    """
    슬라이드 내의 도형들을 스캔하여
    한글 제목/본문, 영어 제목/본문 텍스트 프레임의 shape 인덱스를 자동으로 찾아 매핑한다.
    """
    mapping = {
        'kor_title': None,
        'kor_body':  None,
        'eng_title': None,
        'eng_body':  None,
    }

    for idx, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        raw_text = shape.text_frame.text.strip()
        name_lower = shape.name.strip().lower()

        # 1단계: 텍스트 태그 매칭
        for key, patterns in _TAG_PATTERNS.items():
            if mapping[key] is not None:
                continue
            for pat in patterns:
                if re.search(pat, raw_text, re.IGNORECASE):
                    mapping[key] = idx
                    break

        # 2단계: 도형 이름 매칭
        for key, keywords in _NAME_KEYWORDS.items():
            if mapping[key] is not None:
                continue
            for kw in keywords:
                if kw in name_lower:
                    mapping[key] = idx
                    break

    # 3단계: 기본 템플릿 호환용 폴백 (레거시 인덱스)
    num_shapes = len(slide.shapes)
    if mapping['kor_title'] is None and num_shapes > 1 and slide.shapes[1].has_text_frame:
        mapping['kor_title'] = 1
    if mapping['kor_body'] is None and num_shapes > 6 and slide.shapes[6].has_text_frame:
        mapping['kor_body'] = 6
    if mapping['eng_title'] is None and num_shapes > 5 and slide.shapes[5].has_text_frame:
        mapping['eng_title'] = 5
    if mapping['eng_body'] is None and num_shapes > 7 and slide.shapes[7].has_text_frame:
        mapping['eng_body'] = 7

    return mapping


# ─── 슬라이드 채우기 ──────────────────────────────────────────────────────────
RESPONSIVE_BOLD_PREFIXES = (
    '(회중)', '[회중]', '<회중>',
    '(성도)', '[성도]', '(교인)', '[교인]',
    '(다함께)', '[다함께]', '<다함께>',
    '(다같이)', '[다같이]', '(함께)', '[함께]'
)

RESPONSIVE_ALL_PREFIXES = (
    '(인도)', '[인도]', '<인도>',
    '(회중)', '[회중]', '<회중>',
    '(다함께)', '[다함께]', '<다함께>',
    '(성도)', '[성도]', '(교인)', '[교인]',
    '(다같이)', '[다같이]', '(함께)', '[함께]'
)


def _fill_slide(slide, address, verse, emphases,
                title_shape_idx, body_shape_idx,
                title_style, body_style, bold_font,
                title_use_run=True):
    """슬라이드 하나에 제목(주소)과 본문을 채운다."""
    if title_shape_idx is None and body_shape_idx is None:
        return

    addr_lines = address.split('\n') if address else []
    body_lines = verse.split('\n') if verse else []
    is_multi   = len(addr_lines) > 2

    # 제목 텍스트프레임
    if title_shape_idx is not None and title_shape_idx < len(slide.shapes):
        shape = slide.shapes[title_shape_idx]
        if shape.has_text_frame:
            tf = shape.text_frame
            if not address or not address.strip():
                tf.clear()
            elif title_use_run:
                title_words = addr_lines if is_multi else [w for ln in addr_lines for w in ln.split() if w]
                _set_text_lines(tf, title_words, use_run=True, **title_style)
            else:
                _set_text_lines(tf, addr_lines, use_run=False, **title_style)

    # 본문 텍스트프레임
    if body_shape_idx is not None and body_shape_idx < len(slide.shapes):
        shape = slide.shapes[body_shape_idx]
        if shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.margin_top = Pt(2)
            tf.margin_bottom = Pt(2)
            tf.margin_left = Pt(2)
            tf.margin_right = Pt(2)
            if not verse or not verse.strip():
                pass
            else:
                body_size = body_style['size']
                for i, line in enumerate(body_lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    line_stripped = line.strip()

                    # 교독문 역할 태그 간격 조정
                    is_resp = any(line_stripped.startswith(prefix) for prefix in RESPONSIVE_ALL_PREFIXES)
                    if is_resp:
                        p.space_after = Pt(10)
                        p.line_spacing = 1.2
                    else:
                        p.space_after = Pt(3)
                        p.line_spacing = 1.15

                    line_tokens = line.split()
                    if not line_tokens:
                        full_text = ''
                    elif is_multi:
                        addr_tokens = addr_lines[i].split() if i < len(addr_lines) else []
                        sup_token = addr_tokens[1] if len(addr_tokens) > 1 else ''
                        full_text = f"{_superscript(sup_token)} {_verse_body_text(line)}" if sup_token else line
                    elif line_tokens[0].isdigit():
                        sup_token = line_tokens[0]
                        full_text = f"{_superscript(sup_token)} {_verse_body_text(line)}"
                    else:
                        full_text = line

                    # 회중 또는 다함께 단락인 경우 폰트를 bold_font로 자동 적용
                    para_font = body_style['font']
                    if any(full_text.strip().startswith(prefix) for prefix in RESPONSIVE_BOLD_PREFIXES):
                        para_font = bold_font

                    _write_paragraph_with_emphasis(
                        p, full_text,
                        base_font  = para_font,
                        base_size  = body_size,
                        base_color = body_style['color'],
                        emphases   = emphases,
                        bold_font  = bold_font,
                    )


# ─── PPT 생성 메인 함수 ──────────────────────────────────────────────────────
def add_scripture_to_ppt(
    template_path,
    verse_texts,
    verse_texts_eng,
    style=None,
    bold_font=None,
    output_path="output.pptx",
    return_prs=False
):
    """
    템플릿 PPTX를 기반으로 성경 구절 및 인용문/교독문 슬라이드를 생성한다.

    - template_path   : 템플릿 PPTX 파일 경로
    - verse_texts     : 한글(메인) [(label, verse_text, emphases), ...]
    - verse_texts_eng : 영어(서브) [(label, verse_text, emphases), ...]
    - style           : 사용자 지정 서식 딕셔너리 (생략 시 기본값 사용)
    - bold_font       : '굵게' 서식 및 교독문 회중 파트에 사용할 폰트명
    - output_path     : 생성될 PPTX 저장 경로
    - return_prs      : True이면 저장하지 않고 Presentation 객체를 반환
    """
    resolved_template = resource_path(template_path)
    if not os.path.exists(resolved_template):
        raise FileNotFoundError(f"템플릿 파일을 찾을 수 없습니다: {template_path}")

    prs = Presentation(resolved_template)
    if not prs.slides:
        raise ValueError("템플릿에 슬라이드가 없습니다.")

    # 스타일 딕셔너리 안전 병합
    full_style = get_full_style(style)
    applied_bold_font = (bold_font or '').strip() or DEFAULT_BOLD_FONT

    # 첫 슬라이드 기준 도형 매핑 감지
    mapping = detect_shape_mapping(prs.slides[0])

    # 필요한 수만큼 슬라이드 복제
    while len(prs.slides) < len(verse_texts):
        duplicate_slide_with_blank_layout(prs, prs.slides[-1])

    # 1. 한글(메인 영역) 슬라이드 채우기
    if mapping['kor_body'] is not None or mapping['kor_title'] is not None:
        for idx, (address, verse, emphases) in enumerate(verse_texts):
            _fill_slide(
                slide           = prs.slides[idx],
                address         = address,
                verse           = verse,
                emphases        = emphases,
                title_shape_idx = mapping['kor_title'],
                body_shape_idx  = mapping['kor_body'],
                title_style     = full_style['kor_title'],
                body_style      = full_style['kor_body'],
                bold_font       = applied_bold_font,
                title_use_run   = True,
            )

    # 2. 영어(ESV 영역) 슬라이드 채우기
    if mapping['eng_body'] is not None or mapping['eng_title'] is not None:
        for idx in range(len(prs.slides)):
            if verse_texts_eng and idx < len(verse_texts_eng):
                address, verse, emphases = verse_texts_eng[idx]
            else:
                address, verse, emphases = '', '', []

            _fill_slide(
                slide           = prs.slides[idx],
                address         = address,
                verse           = verse,
                emphases        = emphases,
                title_shape_idx = mapping['eng_title'],
                body_shape_idx  = mapping['eng_body'],
                title_style     = full_style['eng_title'],
                body_style      = full_style['eng_body'],
                bold_font       = applied_bold_font,
                title_use_run   = False,
            )

    if return_prs:
        return prs

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    prs.save(output_path)
    return output_path
